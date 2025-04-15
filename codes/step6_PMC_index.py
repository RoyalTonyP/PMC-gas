import xlrd
from openpyxl import Workbook

# 打开 Excel 文件
wb = xlrd.open_workbook("D:/PyCharm/Py_Projects/gaspaper/results/sub_variables_scores_results.xls")
sheet = wb.sheet_by_index(0)
rows = sheet.nrows
sub_scores = {}

# 遍历每一行数据
for i in range(1, rows):
    file_path = sheet.cell(i, 0).value
    sub_scores[file_path] = {}
    for col in range(1, 10):
        # 获取列名
        col_name = sheet.cell(0, col).value
        # 获取单元格内容
        cell_value = sheet.cell(i, col).value
        # 处理空单元格
        if isinstance(cell_value, str):
            nature = cell_value.replace(' ', '').replace('{', '').replace('}', '').split(',')
            sub_dict = {}
            for ii in nature:
                if ':' in ii:
                    key, value = ii.split(':')
                    try:
                        sub_dict[key] = int(value)
                    except ValueError:
                        print(f"无法将 {value} 转换为整数，跳过该数据。")
            sub_scores[file_path][col_name] = sub_dict
        else:
            sub_scores[file_path][col_name] = {}

# 计算主变量评分和 PMC 指数
for file_path in sub_scores.keys():
    pmc_index = 0
    for main_variable in sub_scores[file_path].keys():
        sub_vars = sub_scores[file_path][main_variable]
        if sub_vars:
            mv_score = sum(sub_vars.values()) / len(sub_vars)
            sub_scores[file_path][main_variable]['主变量评分'] = mv_score
            pmc_index += mv_score
        else:
            sub_scores[file_path][main_variable]['主变量评分'] = 0
    sub_scores[file_path]['PMC指数'] = pmc_index

# 创建新的 Excel 文件
wb = Workbook()
ws = wb.active

# 写入表头
ws.cell(row=1, column=1, value="主变量")

# 写入第一列（主变量名称）
row_0 = 2
main_variables = list(sub_scores[list(sub_scores.keys())[0]].keys())
main_variables.remove('PMC指数')
for main_variable in main_variables:
    ws.cell(row=row_0, column=1, value=main_variable)
    row_0 += 1
ws.cell(row=row_0, column=1, value='PMC指数')

# 写入数据
column = 2
for file_path in sub_scores.keys():
    ws.cell(row=1, column=column, value=file_path)
    row_temp = 2
    for main_variable in main_variables:
        ws.cell(row=row_temp, column=column, value=sub_scores[file_path][main_variable]['主变量评分'])
        row_temp += 1
    ws.cell(row=row_temp, column=column, value=sub_scores[file_path]['PMC指数'])
    column += 1

# 保存新的 Excel 文件
wb.save("D:/PyCharm/Py_Projects/gaspaper/results/main_variable_scores_&_PMC_index.xlsx")