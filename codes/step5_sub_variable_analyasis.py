import xlrd
from openpyxl import Workbook

wb = xlrd.open_workbook("D:/PMC-gas/results/sub_variables_scores_results.xls")
sheet = wb.sheet_by_index(0)
rows = sheet.nrows
sub_scores = {}

for i in range(1, rows):
    file_path = sheet.cell(i, 0).value
    if file_path not in sub_scores:
        sub_scores[file_path] = {}

    # 资源勘探与开发
    timeliness = sheet.cell(i, 1).value.replace(' ', '').replace('{', '').replace('}', '').split(',')
    exploration_dict = {}
    for ii in timeliness:
        parts = ii.split(':')
        if len(parts) == 2:
            try:
                exploration_dict[parts[0]] = int(parts[1])
            except ValueError:
                continue
    sub_scores[file_path]['资源勘探与开发'] = exploration_dict

    # 管网建设与运营
    tools = sheet.cell(i, 2).value.replace(' ', '').replace('{', '').replace('}', '').split(',')
    pipeline_dict = {}
    for ii in tools:
        parts = ii.split(':')
        if len(parts) == 2:
            try:
                pipeline_dict[parts[0]] = int(parts[1])
            except ValueError:
                continue
    sub_scores[file_path]['管网建设与运营'] = pipeline_dict

    # 储气调峰
    function = sheet.cell(i, 3).value.replace(' ', '').replace('{', '').replace('}', '').split(',')
    peak_dict = {}
    for ii in function:
        parts = ii.split(':')
        if len(parts) == 2:
            try:
                peak_dict[parts[0]] = int(parts[1])
            except ValueError:
                continue
    sub_scores[file_path]['储气调峰'] = peak_dict

    # 消费引导
    implementation_agency = sheet.cell(i, 4).value.replace(' ', '').replace('{', '').replace('}', '').split(',')
    consumption_dict = {}
    for ii in implementation_agency:
        parts = ii.split(':')
        if len(parts) == 2:
            try:
                consumption_dict[parts[0]] = int(parts[1])
            except ValueError:
                continue
    sub_scores[file_path]['消费引导'] = consumption_dict

    # 市场准入
    area = sheet.cell(i, 5).value.replace(' ', '').replace('{', '').replace('}', '').split(',')
    market_dict = {}
    for ii in area:
        parts = ii.split(':')
        if len(parts) == 2:
            try:
                market_dict[parts[0]] = int(parts[1])
            except ValueError:
                continue
    sub_scores[file_path]['市场准入资质'] = market_dict

    # 安全环保
    release_agency = sheet.cell(i, 6).value.replace(' ', '').replace('{', '').replace('}', '').split(',')
    environment_dict = {}
    for ii in release_agency:
        parts = ii.split(':')
        if len(parts) == 2:
            try:
                environment_dict[parts[0]] = int(parts[1])
            except ValueError:
                continue
    sub_scores[file_path]['安全环保'] = environment_dict

    # 价格调控
    nature = sheet.cell(i, 7).value.replace(' ', '').replace('{', '').replace('}', '').split(',')
    price_dict = {}
    for ii in nature:
        parts = ii.split(':')
        if len(parts) == 2:
            try:
                price_dict[parts[0]] = int(parts[1])
            except ValueError:
                continue
    sub_scores[file_path]['价格调控'] = price_dict

    # 税收
    measures = sheet.cell(i, 8).value.replace(' ', '').replace('{', '').replace('}', '').split(',')
    tax_dict = {}
    for ii in measures:
        parts = ii.split(':')
        if len(parts) == 2:
            try:
                tax_dict[parts[0]] = int(parts[1])
            except ValueError:
                continue
    sub_scores[file_path]['税收'] = tax_dict

    # 监管
    coverage = sheet.cell(i, 9).value.replace(' ', '').replace('{', '').replace('}', '').split(',')
    supervision_dict = {}
    for ii in coverage:
        parts = ii.split(':')
        if len(parts) == 2:
            try:
                supervision_dict[parts[0]] = int(parts[1])
            except ValueError:
                continue
    sub_scores[file_path]['监管'] = supervision_dict

# print(sub_scores.keys())
# print(sub_scores[list(sub_scores.keys())[0]])

wb = Workbook()
ws = wb.active

# 写入表头
ws['A1'] = "主变量"
ws['B1'] = "子变量"
column = 3
for i in sub_scores.keys():
    ws.cell(row=1, column=column, value=i)

    # 写入数值
    row_temp = 2
    for j in sub_scores[i].keys():
        for k in sub_scores[i][j].keys():
            ws.cell(row=row_temp, column=column, value=sub_scores[i][j][k])
            row_temp += 1
    column += 1

# 写入前两列
row2 = 2
for i in sub_scores[list(sub_scores.keys())[0]].keys():
    ws.cell(row=row2, column=1, value=i)
    for j in sub_scores[list(sub_scores.keys())[0]][i].keys():
        ws.cell(row=row2, column=2, value=j)
        row2 += 1

wb.save("D:/PMC-gas/results/sub_variable_scores_standard.xlsx")
